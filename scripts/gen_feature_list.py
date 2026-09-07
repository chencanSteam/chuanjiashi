# -*- coding: utf-8 -*-
"""生成《传家世 MVP 功能清单》Excel 版（基于当前实现）"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PRIMARY = '1B5E4B'
LIGHT = 'E8F3EF'

wb = Workbook()
ws = wb.active
ws.title = 'MVP功能清单'

thin = Side(style='thin', color='D9D9D9')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
module_font = Font(name='微软雅黑', size=10.5, bold=True, color=PRIMARY)
body_font = Font(name='微软雅黑', size=10.5)
header_fill = PatternFill('solid', fgColor=PRIMARY)
module_fill = PatternFill('solid', fgColor=LIGHT)
wrap = Alignment(vertical='center', wrap_text=True)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

ws.merge_cells('A1:F1')
c = ws['A1']
c.value = '传家世 · AI数字人生与家风传承平台  MVP 功能清单（V1.0）'
c.font = Font(name='微软雅黑', size=14, bold=True, color=PRIMARY)
c.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 32

ws.merge_cells('A2:F2')
c = ws['A2']
c.value = 'MVP 聚焦核心链路「AI 采访 → 传记生成 → 人生档案」，完整版功能（家庭空间/数字人/商城/导出等）不包含在内'
c.font = Font(name='微软雅黑', size=10, color='6B7280')
c.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 22

headers = ['端', '模块', '功能点', '功能说明', '优先级', '状态']
for i, h in enumerate(headers, 1):
    cell = ws.cell(row=3, column=i, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border
ws.row_dimensions[3].height = 24

rows = [
    ('用户端', '账号与登录', [
        ('手机号验证码登录', '获取验证码（60秒重发倒计时）→ 验证登录，未取码/错码拦截，需勾选用户协议', 'P0'),
        ('账号注册', '独立注册页，仅手机号+短信验证码注册，注册成功进入新手引导', 'P0'),
        ('微信授权登录', '微信快捷登录入口（演示为模拟授权）', 'P1'),
        ('账号姓名规则', '展示姓名优先级：实名姓名 > 昵称 > 默认数字名/微信名', 'P0'),
    ]),
    ('用户端', '首页', [
        ('开始智能采访', '主入口，新建档案（基础信息+人生标签）后进入采访', 'P0'),
        ('协作邀请处理', '收到采访协作/人物关系邀请时，首页卡片「同意/拒绝」，同意后建立协作关系', 'P0'),
        ('我协助的传记', '列出可协助的传记，点击进入协助采访', 'P0'),
        ('待办与动态', '待办事项（继续采访等）与最近动态', 'P2'),
    ]),
    ('用户端', 'AI智能采访', [
        ('对话式采访', '三栏布局：左侧主题 / 中间采访区（AI头像+当前问题+回答框）/ 右侧对话记录（自动保存）', 'P0'),
        ('主题采访', '按后台配置的主题进行；无题库，AI 根据主题名称生成开场问题', 'P0'),
        ('AI 追问与衍生', '回答后 AI 自动生成延伸问题并接着问；每题追问≤3次（后台隐性规则，用户无感），可衍生新问题', 'P0'),
        ('语音实时转写', '麦克风录制时输入框切换为录音态（波形+计时），Web Speech API 实时转写为文字', 'P1'),
        ('自定义主题', '用户可新增/删除主题（含预设主题），AI 按主题名出题，删除按档案持久化', 'P1'),
        ('选择传记', '头部切换回答哪本传记：自己的为本人，其余一律为协助', 'P0'),
        ('协作身份判定', '传主=本人（主回答、推进度），其余账号=协助者（补充素材、独立问题集与进度）', 'P0'),
        ('协作邀请（同意制）', '输入对方手机号/身份证号查找账号（回车可查）→ 选关系 → 发送邀请，对方同意后成为协作者，可撤销待同意邀请', 'P0'),
        ('主题内查看协助者', '某主题下有协助回答时可切换到该协助者视角，逐条查看问答（只读）', 'P1'),
        ('作废协助回答', '本人可作废协助者的回答，作废后不作为传记参考，可恢复', 'P1'),
    ]),
    ('用户端', 'AI传记生成', [
        ('选择传记', '头部切换为哪份档案生成传记（帮别人建档即可帮别人生成）', 'P0'),
        ('文风与字数档位', '朴实自然/温情叙事/典雅文言/新闻纪实；短篇约5千字/标准约1.5万字/长篇约3万字', 'P0'),
        ('分章生成与编辑', '章节目录（前言/童年记忆/求学岁月/工作经历/创业之路/家庭生活/人生感悟/后记），支持生成本章、重新生成、润色、插图、手动编辑', 'P0'),
        ('低材料提示', '首次生成且资料完整度<40%时提示，可继续采访或仍然生成，不拦截', 'P1'),
        ('导入已有传记', '粘贴文本或上传文件，按章节标题自动拆分', 'P2'),
        ('保存到我的传记', '生成稿保存至我的传记', 'P0'),
    ]),
    ('用户端', '人生档案', [
        ('人生时间轴', '按年份记录人生大事；编辑事件（年份下拉、标题、描述、附件照片）', 'P0'),
        ('多媒体档案库', '照片/视频/音频/文档素材管理与预览', 'P1'),
        ('人物关系图谱', '以传主为中心展示关系；维护关系=手机号/身份证查找+对方同意后写入图谱', 'P0'),
        ('地点足迹', '人生地点地图与足迹列表', 'P2'),
        ('成就与作品', '荣誉、作品记录', 'P2'),
        ('隐私与权限', '角色仅档案所有者/观察者；5个模块（时间轴/多媒体/关系图谱/地点足迹/成就作品）各自设置家人可见/公开展示/仅自己', 'P0'),
        ('新建档案与人生标签', '新建档案可选预设/自定义人生标签，标签在档案概览展示并驱动采访主题', 'P1'),
        ('档案概览', '固定右侧栏，切换标签不变：头像、基本信息、人生标签、档案动态', 'P1'),
    ]),
    ('用户端', '个人中心', [
        ('实名认证', '真实姓名+18位身份证号（格式校验），认证后账号姓名切换为实名', 'P0'),
        ('绑定微信', '绑定/解绑微信', 'P1'),
        ('绑定手机号', '更换手机号需短信验证码校验', 'P1'),
        ('退出登录', '账号退出', 'P0'),
    ]),
    ('运营后台', '用户管理', [
        ('用户列表', '昵称（系统默认用户+尾号）、手机号、注册时间、实名状态、状态；搜索与筛选', 'P0'),
        ('禁用/启用', '账号状态控制，禁用后无法登录使用', 'P0'),
        ('实名状态展示', '实名认证为系统自动完成，后台只读展示（未认证/认证中/已实名/认证失败）', 'P0'),
        ('用户详情', '基础信息、推广关系、佣金明细（抽屉）', 'P2'),
    ]),
    ('运营后台', '人物档案管理', [
        ('档案列表', '主人姓名、创建人、素材数（图/音/文）、隐私状态（只读）、完整度、创建时间', 'P0'),
        ('档案详情', '基础信息与素材统计（抽屉），隐私状态由用户自管，后台不可改', 'P1'),
    ]),
    ('运营后台', 'AI任务管理', [
        ('任务队列', '按用户（昵称+手机号）展示任务：类型、状态、Token消耗、耗时、创建时间，失败可重试', 'P0'),
        ('Token 成本', '总消耗/任务数/估算成本；按类型汇总 + 按用户汇总（消耗排序）', 'P1'),
        ('模板管理', '提示词/问卷/文风/规则模板：启停开关 + 编辑名称与内容', 'P1'),
        ('采访主题', '主题增删改、启停、排序自动重排；预设主题6个，新增主题 AI 出题；配置实时生效到采访页', 'P0'),
    ]),
    ('运营后台', '角色权限', [
        ('角色权限说明', '通用角色权限管理模块说明页；MVP 阶段单一管理员角色，多角色后续开放', 'P2'),
    ]),
]

side_counts = {}
for side, module, items in rows:
    side_counts[side] = side_counts.get(side, 0) + len(items) + 1

current_row = 4
side_start = {}
for side, module, items in rows:
    ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=6)
    ws.cell(row=current_row, column=1, value=side)
    ws.cell(row=current_row, column=2, value=module)
    for col in range(1, 7):
        cell = ws.cell(row=current_row, column=col)
        cell.fill = module_fill
        cell.font = module_font
        cell.border = border
        cell.alignment = wrap
    if side not in side_start:
        side_start[side] = current_row
    current_row += 1
    for feat, desc, prio in items:
        ws.cell(row=current_row, column=3, value=feat)
        ws.cell(row=current_row, column=4, value=desc)
        ws.cell(row=current_row, column=5, value=prio)
        ws.cell(row=current_row, column=6, value='已实现')
        for col in range(1, 7):
            cell = ws.cell(row=current_row, column=col)
            cell.font = body_font
            cell.border = border
            cell.alignment = center if col in (1, 5, 6) else wrap
        current_row += 1

for side, start in side_start.items():
    ws.merge_cells(start_row=start, start_column=1, end_row=start + side_counts[side] - 1, end_column=1)
    ws.cell(row=start, column=1).alignment = center

# Sheet2: 不包含的功能
ws2 = wb.create_sheet('MVP不包含（完整版规划）')
ws2.merge_cells('A1:C1')
c = ws2['A1']
c.value = 'MVP 不包含的功能（完整版规划）'
c.font = Font(name='微软雅黑', size=13, bold=True, color=PRIMARY)
ws2.row_dimensions[1].height = 28
for i, h in enumerate(['方向', '功能', '说明'], 1):
    cell = ws2.cell(row=2, column=i, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border

excluded = [
    ('家庭协作', '家庭空间 / 数字家谱 / AI家风馆', '家庭共享、家谱编辑、家风内容与活动'),
    ('数字永生', '数字人 / 数字陪伴', '数字人格训练、语音对话、训练记录报告'),
    ('商业化', '传家商城 / 拼团 / 传记师服务 / 合伙人', '实体书、衍生品、服务市场与分润'),
    ('传记衍生', '导出Word/PDF / 实体书排版 / 衍生内容 / 快捷生成', '导出与实体书、人生金句/家风总结等提炼'),
    ('内容生态', '公开传记书架 / 数字博物馆 / 老照片修复', '公开内容与展馆、照片修复工具'),
    ('政务服务', '政务服务 / 继承方案', '政务对接与数字资产继承'),
    ('运营后台扩展', '总览看板 / 业务管理 / 财务分润 / 内容合规 / 消息通知 / 二维码管理', '完整后台各业务模块'),
]
rr = 3
for direction, feat, desc in excluded:
    ws2.cell(row=rr, column=1, value=direction)
    ws2.cell(row=rr, column=2, value=feat)
    ws2.cell(row=rr, column=3, value=desc)
    for col in range(1, 4):
        cell = ws2.cell(row=rr, column=col)
        cell.font = body_font
        cell.border = border
        cell.alignment = wrap
    rr += 1

widths = [10, 14, 22, 62, 8, 8]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for i, w in enumerate([12, 40, 40], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = 'A4'
ws2.freeze_panes = 'A3'

out = '../传家世MVP功能清单_最新版.xlsx'
wb.save(out)
print('saved:', out)
